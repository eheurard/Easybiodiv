import openpyxl
from django.test import TestCase

from dashboard.models import (
    FLOW_RULES, REVENUE_PRODUCTION_ONLY_MESSAGE, SUPPLY_SELF_LOOP_MESSAGE, Asset, Commodity,
    Company, CommodityModifier, Country, Flow, FlowKind,
)
from imports.services.excel_export import build_flow_export
from imports.services.excel_parser import parse_file
from imports.services.excel_template import build_template
from imports.services.importer import save_import
from imports.tests.test_excel_parser import _make_xlsx, _sheet


class FlowSheetTests(TestCase):

    def setUp(self):
        self.brazil = Country.objects.create(
            name='Brésil', water_ownership='pub', land_ownership='priv')
        self.mine = Asset.objects.create(
            name='Mine', latitude=1.0, longitude=2.0, country=self.brazil)
        self.acme = Company.objects.create(name='Acme')
        Commodity.objects.create(name='Bœuf')

    def _row(self, **values):
        row = {
            'kind': 'SUPPLY', 'what': 'Bœuf', 'from_type': 'country', 'from_name': 'Brésil',
            'to_type': 'company', 'to_name': 'Acme', 'year': '2026', 'quantity': '1000',
        }
        row.update(values)
        return row

    def _parse(self, *rows):
        return parse_file(_make_xlsx({'Flow': _sheet('Flow', *rows)}))

    def test_one_row_of_each_kind_is_imported(self):
        parsed = self._parse(
            self._row(),
            self._row(kind='PRODUCTION', from_type='asset', from_name='Mine',
                      to_type='', to_name='', estimated_revenue='500'),
            self._row(kind='CONSUMPTION', what='water', from_type='milieu', from_name='',
                      to_type='asset', to_name='Mine'),
            self._row(kind='EMISSION', what='co2', scope='scope 1', from_type='company',
                      from_name='Acme', to_type='milieu', to_name=''),
            self._row(kind='WASTE', what='waste', from_type='asset', from_name='Mine',
                      to_type='', to_name=''),
        )
        self.assertEqual([r['status'] for r in parsed['Flow']], ['ok'] * 5)
        self.assertEqual(save_import(parsed)['Flow'], 5)
        supply = Flow.objects.get(kind=FlowKind.SUPPLY)
        self.assertEqual(
            (supply.from_country, supply.to_company, supply.quantity),
            (self.brazil, self.acme, 1000.0))
        production = Flow.objects.get(kind=FlowKind.PRODUCTION)
        self.assertEqual((production.from_asset, production.estimated_revenue), (self.mine, 500.0))
        consumption = Flow.objects.get(kind=FlowKind.CONSUMPTION)
        self.assertEqual((consumption.from_environment, consumption.what.key), (True, 'water'))
        emission = Flow.objects.get(kind=FlowKind.EMISSION)
        self.assertEqual((emission.scope, emission.to_environment), ('Scope 1', True))

    def test_lowercase_kind_is_accepted(self):
        self.assertEqual(self._parse(self._row(kind='supply'))['Flow'][0]['status'], 'ok')

    def test_unknown_commodity_is_error(self):
        row = self._parse(self._row(what='Licorne'))['Flow'][0]
        self.assertEqual(row['status'], 'error')
        self.assertIn('what', row['message'])

    def test_name_next_to_milieu_is_error(self):
        row = self._parse(self._row(
            kind='EMISSION', what='co2', from_type='company', from_name='Acme',
            to_type='milieu', to_name='Océan'))['Flow'][0]
        self.assertEqual(row['status'], 'error')
        self.assertIn('to_name', row['message'])

    def test_unknown_endpoint_name_is_error(self):
        row = self._parse(self._row(from_name='Atlantide'))['Flow'][0]
        self.assertEqual(row['status'], 'error')
        self.assertIn('from_name', row['message'])

    def test_rule_violation_uses_the_model_message(self):
        row = self._parse(self._row(
            kind='EMISSION', what='co2', from_type='company', from_name='Acme',
            to_type='', to_name=''))['Flow'][0]
        self.assertEqual(row['status'], 'error')
        self.assertEqual(row['message'], FLOW_RULES['EMISSION'].message)

    def test_revenue_outside_production_is_error(self):
        row = self._parse(self._row(estimated_revenue='10'))['Flow'][0]
        self.assertEqual(row['message'], REVENUE_PRODUCTION_ONLY_MESSAGE)

    def test_existing_flow_is_duplicate_even_when_what_is_given_by_key(self):
        Flow.objects.create(
            kind=FlowKind.EMISSION, what=Commodity.objects.technical('co2'),
            from_company=self.acme, to_environment=True, year=2024, quantity=1.0)
        row = self._parse(self._row(
            kind='EMISSION', what='CO2', from_type='company', from_name='Acme',
            to_type='milieu', to_name='', year='2024'))['Flow'][0]
        self.assertEqual(row['status'], 'duplicate')

    def test_removed_sheet_is_reported(self):
        parsed = parse_file(_make_xlsx({'Production': [['asset_name'], ['Mine']]}))
        self.assertEqual(parsed['Production'][0]['status'], 'error')
        self.assertIn('Flow', parsed['Production'][0]['message'])

    def test_reference_sheet_lists_units_keys_and_endpoint_types(self):
        ws = openpyxl.load_workbook(build_template())['_Référence']
        values = {cell.value for row in ws.iter_rows() for cell in row}
        self.assertIn('CO₂ — tCO₂e — co2', values)
        self.assertIn('milieu', values)

    def test_tier_out_of_range_is_clamped(self):
        """Ruling tâche 8 : `_tier` clampe la cellule sur l'intervalle 0-3 du modèle."""
        parsed = self._parse(self._row(tier='9'))
        save_import(parsed)
        self.assertEqual(Flow.objects.get().tier, 3)

    def test_comma_decimal_quantity_is_imported(self):
        parsed = self._parse(self._row(quantity='12,5'))
        self.assertEqual(parsed['Flow'][0]['status'], 'ok')
        save_import(parsed)
        self.assertEqual(Flow.objects.get().quantity, 12.5)

    def test_invalid_year_is_error(self):
        row = self._parse(self._row(year='FY2024'))['Flow'][0]
        self.assertEqual(row['status'], 'error')
        self.assertIn('year', row['message'])

    def test_invalid_quantity_is_error(self):
        row = self._parse(self._row(quantity='abc'))['Flow'][0]
        self.assertEqual(row['status'], 'error')
        self.assertIn('quantity', row['message'])

    def test_invalid_tier_is_error(self):
        row = self._parse(self._row(
            kind='PRODUCTION', from_type='asset', from_name='Mine', to_type='', to_name='',
            tier='deux'))['Flow'][0]
        self.assertEqual(row['status'], 'error')
        self.assertIn('tier', row['message'])

    def test_invalid_estimated_revenue_is_error(self):
        row = self._parse(self._row(
            kind='PRODUCTION', from_type='asset', from_name='Mine', to_type='', to_name='',
            estimated_revenue='n/a'))['Flow'][0]
        self.assertEqual(row['status'], 'error')
        self.assertIn('estimated_revenue', row['message'])

    def test_supply_from_the_global_market_is_ok(self):
        row = self._parse(self._row(from_type='', from_name=''))['Flow'][0]
        self.assertEqual(row['status'], 'ok', row.get('message'))

    def test_supply_self_loop_is_error(self):
        row = self._parse(self._row(
            from_type='asset', from_name='Mine', to_type='asset', to_name='Mine'))['Flow'][0]
        self.assertEqual(row['status'], 'error')
        self.assertEqual(row['message'], SUPPLY_SELF_LOOP_MESSAGE)

    def test_duplicate_row_within_the_same_file_is_duplicate(self):
        rows = self._parse(self._row(), self._row())['Flow']
        self.assertEqual([r['status'] for r in rows], ['ok', 'duplicate'])

    def test_uppercase_endpoint_type_is_accepted(self):
        row = self._parse(self._row(from_type='ASSET', from_name='Mine'))['Flow'][0]
        self.assertEqual(row['status'], 'ok')


class FlowModifiersSheetTests(TestCase):

    def setUp(self):
        self.brazil = Country.objects.create(
            name='Brésil', water_ownership='pub', land_ownership='priv')
        self.mine = Asset.objects.create(
            name='Ferme', latitude=1.0, longitude=2.0, country=self.brazil)
        self.soy = Commodity.objects.create(name='Soja')
        self.rtrs = CommodityModifier.objects.create(commodity=self.soy, name='Certifié RTRS')
        self.bio = CommodityModifier.objects.create(commodity=self.soy, name='Bio')

    def _row(self, **values):
        row = {
            'kind': 'PRODUCTION', 'what': 'Soja', 'from_type': 'asset', 'from_name': 'Ferme',
            'year': '2026', 'quantity': '100',
        }
        row.update(values)
        return row

    def _parse(self, *rows):
        return parse_file(_make_xlsx({'Flow': _sheet('Flow', *rows)}))

    def test_flow_modifiers_roundtrip(self):
        flow = Flow.objects.create(
            kind=FlowKind.PRODUCTION, what=self.soy, from_asset=self.mine, year=2026,
            quantity=100.0)
        flow.modifiers.set([self.rtrs, self.bio])
        buffer, count = build_flow_export()
        self.assertEqual(count, 1)
        sheet = openpyxl.load_workbook(buffer)['Flow']
        header = [c.value for c in sheet[1]]
        self.assertEqual(header[-1], 'modifiers')
        self.assertEqual(sheet.cell(2, len(header)).value, 'Bio; Certifié RTRS')
        buffer.seek(0)
        flow.delete()
        parsed = parse_file(buffer)
        self.assertEqual([r['status'] for r in parsed['Flow']], ['ok'])
        save_import(parsed)
        self.assertEqual(
            set(Flow.objects.get().modifiers.all()), {self.rtrs, self.bio})

    def test_modifier_cell_ignores_case_and_spaces(self):
        parsed = self._parse(self._row(modifiers=' bio ;CERTIFIÉ rtrs'))
        self.assertEqual(parsed['Flow'][0]['status'], 'ok')
        save_import(parsed)
        self.assertEqual(set(Flow.objects.get().modifiers.all()), {self.rtrs, self.bio})

    def test_unknown_modifier_rejected(self):
        row = self._parse(self._row(modifiers='Bio; Fantôme'))['Flow'][0]
        self.assertEqual(row['status'], 'error')
        self.assertEqual(
            row['message'],
            'Modificateur introuvable pour la commodité « Soja » : « Fantôme »')

    def test_modifier_of_another_commodity_rejected(self):
        wheat = Commodity.objects.create(name='Blé')
        CommodityModifier.objects.create(commodity=wheat, name='Blé bio')
        row = self._parse(self._row(modifiers='Blé bio'))['Flow'][0]
        self.assertEqual(row['status'], 'error')

    def test_certified_and_classic_rows_not_duplicates(self):
        parsed = self._parse(
            self._row(modifiers='Certifié RTRS'),
            self._row(modifiers=''),
        )
        self.assertEqual([r['status'] for r in parsed['Flow']], ['ok', 'ok'])
        self.assertEqual(save_import(parsed)['Flow'], 2)
        self.assertEqual(Flow.objects.count(), 2)
        again = self._parse(self._row(modifiers='Certifié RTRS'), self._row(modifiers=''))
        self.assertEqual([r['status'] for r in again['Flow']], ['duplicate', 'duplicate'])
