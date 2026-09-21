import io
import json
import os
import openpyxl
from django.test import TestCase, override_settings
from django.urls import reverse
from django.core.files.uploadedfile import SimpleUploadedFile
from authentication.models import User

import tempfile


def _make_minimal_xlsx():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet('Country')
    ws.append(['name', 'water_ownership', 'land_ownership', 'water_Governance', 'land_Governance'])
    ws.append(['TestLand', 'pub', 'priv', '', ''])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@override_settings(MEDIA_ROOT=tempfile.mkdtemp())
class ViewsAccessTest(TestCase):
    def setUp(self):
        self.creator = User.objects.create_user('creator', password='pass', role='CREATOR')
        self.subscriber = User.objects.create_user('sub', password='pass', role='SUBSCRIBER')

    def test_index_anonymous_redirects(self):
        resp = self.client.get(reverse('imports:index'))
        self.assertEqual(resp.status_code, 302)
        self.assertIn('/auth/login/', resp['Location'])

    def test_index_subscriber_403(self):
        self.client.login(username='sub', password='pass')
        resp = self.client.get(reverse('imports:index'))
        self.assertEqual(resp.status_code, 403)

    def test_index_creator_200(self):
        self.client.login(username='creator', password='pass')
        resp = self.client.get(reverse('imports:index'))
        self.assertEqual(resp.status_code, 200)

    def test_download_template_returns_xlsx(self):
        self.client.login(username='creator', password='pass')
        resp = self.client.get(reverse('imports:download_template'))
        self.assertEqual(resp.status_code, 200)
        self.assertIn('spreadsheetml', resp['Content-Type'])

    def test_upload_bad_extension_shows_error(self):
        self.client.login(username='creator', password='pass')
        f = SimpleUploadedFile('test.csv', b'name\nFrance\n', content_type='text/csv')
        resp = self.client.post(reverse('imports:upload'), {'excel_file': f})
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'Format invalide')

    def test_upload_valid_file_redirects_to_preview(self):
        self.client.login(username='creator', password='pass')
        xlsx_bytes = _make_minimal_xlsx()
        f = SimpleUploadedFile('data.xlsx', xlsx_bytes,
                               content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        resp = self.client.post(reverse('imports:upload'), {'excel_file': f})
        self.assertRedirects(resp, reverse('imports:preview'), fetch_redirect_response=False)
        self.assertIn('import_key', self.client.session)

    def test_preview_renders_data_rows(self):
        self.client.login(username='creator', password='pass')
        xlsx_bytes = _make_minimal_xlsx()
        f = SimpleUploadedFile('data.xlsx', xlsx_bytes,
                               content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.client.post(reverse('imports:upload'), {'excel_file': f})
        resp = self.client.get(reverse('imports:preview'))
        self.assertEqual(resp.status_code, 200)
        # The actual data row must be rendered, not the empty-tab notice.
        self.assertContains(resp, 'TestLand')
        self.assertContains(resp, 'badge--ok')
        self.assertNotContains(resp, 'Aucune ligne dans cet onglet')

    def test_preview_without_session_redirects(self):
        self.client.login(username='creator', password='pass')
        resp = self.client.get(reverse('imports:preview'))
        self.assertRedirects(resp, reverse('imports:index'), fetch_redirect_response=False)

    def test_full_workflow_upload_preview_confirm(self):
        self.client.login(username='creator', password='pass')
        xlsx_bytes = _make_minimal_xlsx()
        f = SimpleUploadedFile('data.xlsx', xlsx_bytes,
                               content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.client.post(reverse('imports:upload'), {'excel_file': f})
        resp = self.client.get(reverse('imports:preview'))
        self.assertEqual(resp.status_code, 200)
        resp = self.client.post(reverse('imports:confirm'))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'importé')
        from dashboard.models import Country
        self.assertTrue(Country.objects.filter(name='TestLand').exists())


@override_settings(MEDIA_ROOT=tempfile.mkdtemp())
class ExportAllViewTests(TestCase):
    """Bouton « Exporter toutes les données » : réservé aux créateurs, journalisé."""

    def setUp(self):
        self.creator = User.objects.create_user('creator', password='pass', role='CREATOR')
        User.objects.create_user('sub', password='pass', role='SUBSCRIBER')

    def test_anonymous_is_sent_to_login(self):
        resp = self.client.get(reverse('imports:export_all'))
        self.assertEqual(resp.status_code, 302)
        self.assertIn('/auth/login/', resp['Location'])

    def test_subscriber_is_forbidden(self):
        self.client.login(username='sub', password='pass')
        resp = self.client.get(reverse('imports:export_all'))
        self.assertEqual(resp.status_code, 403)

    def test_creator_downloads_a_dated_workbook_with_every_sheet(self):
        from django.utils import timezone

        from imports.services.constants import SHEET_COLUMNS

        self.client.login(username='creator', password='pass')
        resp = self.client.get(reverse('imports:export_all'))
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(
            resp['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        self.assertEqual(
            resp['Content-Disposition'],
            f'attachment; filename="easybiodiv_export_{timezone.localdate():%Y-%m-%d}.xlsx"',
        )
        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        self.assertEqual(wb.sheetnames, list(SHEET_COLUMNS))

    def test_every_export_is_logged_with_its_author(self):
        self.client.login(username='creator', password='pass')
        with self.assertLogs('imports.views', level='INFO') as logs:
            self.client.get(reverse('imports:export_all'))
        self.assertEqual(len(logs.output), 1)
        self.assertIn('Export complet par creator', logs.output[0])

    def test_export_is_post_free(self):
        # Pas de formulaire : un simple lien de téléchargement, en GET seulement.
        self.client.login(username='creator', password='pass')
        resp = self.client.post(reverse('imports:export_all'))
        self.assertEqual(resp.status_code, 405)

    def test_button_is_shown_to_creators_on_the_import_page(self):
        self.client.login(username='creator', password='pass')
        resp = self.client.get(reverse('imports:index'))
        self.assertContains(resp, reverse('imports:export_all'))
        self.assertContains(resp, 'Exporter toutes les données')
